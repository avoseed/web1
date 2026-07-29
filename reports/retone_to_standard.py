# -*- coding: utf-8 -*-
"""외부 덱을 SSG/ZETTA 표준 톤앤매너로 일괄 변환하는 재사용 도구.

사용:  python3 retone_to_standard.py <입력.pptx> <출력.pptx> [--keep-size]
       --keep-size : 판형·좌표·폰트 크기를 원본 유지하고 톤(폰트체·팔레트·표 스타일)만 변경

적용 사례: OGRP/CFC 재고 분석 10장 덱(나눔바른고딕·SSG 레드 E50012·검정 표헤더) 변환.

목표 톤 = reports/ssg-*.pptx 실측(ZETTA v4 표준)
  판형 27.52 × 19.05cm · 맑은 고딕(latin·ea·cs 3중 지정) · 좌·우·하단 1.5cm 액자
  검정 222222 / 크림슨 C30C3E(강조) / 표 헤더 EAEEF6 + 검정 볼드
  강조 셀 DEEBF7 + 002060 / 보조 595959

주의: python-pptx 의 `shape.line.color` · `fill.fore_color` 접근은 없던 <a:ln>/채움을
      생성해 테두리가 새로 그려진다. 따라서 색 치환은 **XML srgbClr 직접 교체**로 처리한다.
"""
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

import sys

SRC = sys.argv[1] if len(sys.argv) > 1 else "src.pptx"
DST = sys.argv[2] if len(sys.argv) > 2 else "out.pptx"
KEEP = "--keep-size" in sys.argv          # 판형 원본 유지 모드
FONT = "맑은 고딕"

CMAP = {                      # 원본 → 표준
    "E50012": "C30C3E",       # SSG 레드 → 크림슨
    "6B7280": "595959",       # 보조 회색 → SUB_GRAY
    "E5E7EB": "D8E0EC",       # 2차 헤더 → TH_SECOND
    "F3F4F6": "F2F2F2",       # 옅은 채움
    "FFF59D": "DEEBF7",       # 노랑 하이라이트 → HL_FILL
}
HDR_BLACK, TH_PRIMARY, HL_TEXT = "222222", "EAEEF6", "002060"
NAVY = "000066"           # 표준: 면 채움 전용(글자색 금지)


def shape_fill_hex(sh):
    """도형 채움색 XML 안전 조회(생성 부작용 없음)."""
    spPr = sh._element.find(qn("p:spPr"))
    if spPr is None:
        return None
    sf = spPr.find(qn("a:solidFill"))
    if sf is None:
        return None
    clr = sf.find(qn("a:srgbClr"))
    return clr.get("val").upper() if clr is not None else None

prs = Presentation(SRC)
NW, NH = Cm(27.52), Cm(19.05)
SRC_L, SRC_R, SRC_T, SRC_B = 1.50, 28.20, 0.85, 20.35     # 원본 콘텐츠 실측(입력 덱에 맞게 조정)
DST_L, DST_R, DST_T, DST_B = 1.50, 26.02, 0.82, 17.55     # SPEC §3-4-0 액자
if KEEP:
    kx = ky = KF = 1.0
else:
    kx = (DST_R - DST_L) / (SRC_R - SRC_L)
    ky = (DST_B - DST_T) / (SRC_B - SRC_T)
    KF = round(ky * 20) / 20
print("kx=%.4f ky=%.4f 폰트=%.2f" % (kx, ky, KF))

mapx = (lambda v: v) if KEEP else (lambda v: Emu(int((DST_L + (v / 360000.0 - SRC_L) * kx) * 360000)))
mapy = (lambda v: v) if KEEP else (lambda v: Emu(int((DST_T + (v / 360000.0 - SRC_T) * ky) * 360000)))


def set_font(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def fix_runs(tf, color=None, bold=None):
    for para in tf.paragraphs:
        for r in para.runs:
            set_font(r)
            if r.font.size is not None:
                r.font.size = Pt(max(7.0, round(r.font.size.pt * KF * 2) / 2))
            if color is not None:
                r.font.color.rgb = RGBColor.from_string(color)
            if bold is not None:
                r.font.bold = bold


def cell_fill_hex(cell):
    """셀 채움색을 XML 로 안전 조회(생성 부작용 없음)."""
    tcPr = cell._tc.find(qn("a:tcPr"))
    if tcPr is None:
        return None
    sf = tcPr.find(qn("a:solidFill"))
    if sf is None:
        return None
    clr = sf.find(qn("a:srgbClr"))
    return clr.get("val").upper() if clr is not None else None


# ── 1) 기하 스케일 + 폰트 + 표 헤더/강조 셀 특수 처리 ─────────────
n_hdr = n_hl = n_navy = 0
for slide in prs.slides:
    for sh in slide.shapes:
        if sh.left is not None:
            sh.left, sh.width = mapx(sh.left), Emu(int(sh.width * kx))
        if sh.top is not None:
            sh.top, sh.height = mapy(sh.top), Emu(int(sh.height * ky))
        if shape_fill_hex(sh) == HDR_BLACK:      # 검정 채움 도형 → 표준 네이비 면
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            n_navy += 1
        if sh.has_text_frame:
            fix_runs(sh.text_frame)
        if sh.has_table:
            t = sh.table
            for c in t.columns:
                c.width = Emu(int(c.width * kx))
            for r_ in t.rows:
                r_.height = Emu(int(r_.height * ky))
            for row in t.rows:
                for cell in row.cells:
                    hexv = cell_fill_hex(cell)
                    if hexv == HDR_BLACK:                    # 검정 헤더 → 연블루+검정 볼드
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(TH_PRIMARY)
                        fix_runs(cell.text_frame, color=HDR_BLACK, bold=True)
                        n_hdr += 1
                    elif hexv == "FFF59D":                   # 노랑 → 강조 셀
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(CMAP["FFF59D"])
                        fix_runs(cell.text_frame, color=HL_TEXT, bold=True)
                        n_hl += 1
                    else:
                        fix_runs(cell.text_frame)

# ── 2) 남은 색은 XML srgbClr 직접 치환 (테두리 생성 부작용 없음) ──
n_col = 0
for slide in prs.slides:
    for el in slide._element.iter(qn("a:srgbClr")):
        v = (el.get("val") or "").upper()
        if v in CMAP:
            el.set("val", CMAP[v])
            n_col += 1

if not KEEP:
    prs.slide_width, prs.slide_height = NW, NH
prs.save(DST)
print("saved %s | 헤더 %d셀 · 강조 %d셀 · 네이비면 %d개 · 색치환 %d건" % (DST, n_hdr, n_hl, n_navy, n_col))
