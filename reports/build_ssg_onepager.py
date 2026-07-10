# -*- coding: utf-8 -*-
"""SSG닷컴 '2시간 배송' 원페이지 보고서 — ZETTA 표준 PPT 빌더(v4.1) 기반.

P계열 본문 1매. 리드메시지 → 도입 배경 → 시범 운영 개요 → 서비스 비교 →
품목 비교·확대 로드맵 순의 논리 구성.
개체 간 수직 간격은 표준 토큰(BODY_TOP_LEAD·GAP_SECTION·GAP_HEAD·LINE_H)으로 고정.
출처: 비즈워치 「[인사이드 스토리] SSG닷컴, '2시간 실험'에 나선 진짜 이유는」('26. 7. 10.)
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "zetta-ppt-standard"))
from zetta_ppt_standard import *
from zetta_ppt_standard import _txt
from pptx.enum.text import PP_ALIGN

SRC = "※ 출처 : 비즈워치 「SSG닷컴, '2시간 실험'에 나선 진짜 이유는」('26. 7. 10. 정혜인 기자)"
HALF_W = 12.30
COL2_L = MARGIN_L + HALF_W + 0.66
SECTION_H = 0.55            # 소제목 텍스트박스 높이


def section(slide, l, t, w, text):
    """검정 볼드 소제목 (□ 섹션 라벨)."""
    return _txt(slide, l, t, w, SECTION_H, "□ " + text, size=13, bold=True, color=BLACK)


def bullet(slide, l, t, w, text, size=10.5):
    return _txt(slide, l + 0.55, t, w - 0.55, LINE_H, "- " + text, size=size, color=INK)


def left_align_body(tbl, cols, header_rows=1):
    for i, row in enumerate(tbl.rows):
        if i < header_rows:
            continue
        for j in cols:
            tbl.cell(i, j).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT


prs = new_deck()
s = add_content(prs, "이마트 퀵커머스", "SSG닷컴 '2시간 배송' 도입 전략", 1, 1,
                tier="P",
                lead="양재·하남점 '2시간 내 배송' 시범 도입 — 대용량 장보기 공략, PP센터 가동률 제고",
                unit="(기준 : '26. 7월)", footnote=SRC)

y = BODY_TOP_LEAD                          # 본문 커서 (리드메시지 사용 시 4.10)

# 1. 도입 배경
section(s, MARGIN_L, y, 25.26, "도입 배경")
bullet(s, MARGIN_L, y + GAP_HEAD, 25.26,
       "이륜차 퀵커머스(바로퀵)는 적재량 한계로 소량 배송에 한정 → 대형마트 주력인 3~4인 가구 대용량 장보기 공백")
bullet(s, MARGIN_L, y + GAP_HEAD + LINE_H, 25.26,
       "바로퀵 6월 매출 1월 대비 197% 증가로 온디맨드 수요 실증 → 1시간(소량·이륜차) / 2시간(대용량·사륜차) 세분화 대응 (B마트도 대용량 확대 중)")
y += GAP_HEAD + 2 * LINE_H + GAP_SECTION

# 2. 시범 운영 개요
section(s, MARGIN_L, y, 25.26, "시범 운영 개요")
bullet(s, MARGIN_L, y + GAP_HEAD, 25.26,
       "오후 8시까지 주문 시 결제 시점 기준 2시간 내 배송 완료, 기존 예약배송 병행 선택 가능")
bullet(s, MARGIN_L, y + GAP_HEAD + LINE_H, 25.26,
       "쓱배송 사륜 차량 활용으로 무거운 대용량 상품까지 즉시 배송 — 점포당 약 15만 종 상품이 기반")
y += GAP_HEAD + 2 * LINE_H + GAP_SECTION

# 3. 배송 서비스 비교
MAIN_TBL_H = 4.50
section(s, MARGIN_L, y, 25.26, "이마트·SSG닷컴 배송 서비스 비교")
data = [
    ["구분", "배송 속도", "운송 수단", "취급 품목", "특징"],
    ["주간배송", "예약배송 (구간당 4~5시간)", "사륜차", "점포 상품", "지역별 2~5개 시간 구간"],
    ["새벽배송", "새벽 시간대", "사륜차", "네오 취급 상품", "물류센터 '네오' 한정 (PP센터 불가)"],
    ["트레이더스배송", "예약배송", "사륜차", "트레이더스 상품", "쓱배송의 한 축"],
    ["바로퀵", "1시간 내외 (반경 약 3km)", "이륜차", "약 1만 종", "소량·1~2인 가구, 6월 매출 1월比 197%↑"],
    ["2시간 배송 (신규)", "2시간 이내 (20시 마감)", "사륜차", "약 15만 종 기반", "대용량 즉시배송"],
]
tbl = add_fin_table(s, MARGIN_L, y + GAP_HEAD, 25.26, MAIN_TBL_H, data,
                    col_w=[3.9, 5.3, 2.5, 3.6, 9.96], header_rows=1, font_size=10)
left_align_body(tbl, [1, 3, 4])
y += GAP_HEAD + MAIN_TBL_H + GAP_SECTION

# 4. 하단 2단 (취급 품목 수 비교 / 확대 로드맵)
SUB_TBL_H = 2.60
section(s, MARGIN_L, y, HALF_W, "취급 품목 수 비교")
data = [
    ["구분", "품목 수", "비고"],
    ["이마트 점포", "약 150,000", "2시간 배송 기반"],
    ["바로퀵", "약 10,000", "이륜차 적재 한계"],
    ["B마트", "약 20,000", "퀵커머스 1위"],
]
tbl = add_fin_table(s, MARGIN_L, y + GAP_HEAD, HALF_W, SUB_TBL_H, data,
                    col_w=[3.5, 3.6, 5.2], header_rows=1,
                    header_fill=TH_SECOND, font_size=10)
left_align_body(tbl, [2])

section(s, COL2_L, y, HALF_W, "확대 로드맵")
data = [
    ["시기", "계획"],
    ["'26. 7월", "양재·하남점 시범 운영 (7. 9.)"],
    ["'26. 8월", "서울 주요 지역 확대"],
    ["'26. 9월 → 연말", "전국 확대 → 50여 개 점포 목표"],
]
tbl = add_fin_table(s, COL2_L, y + GAP_HEAD, HALF_W, SUB_TBL_H, data,
                    col_w=[3.9, 8.4], header_rows=1,
                    header_fill=TH_SECOND, font_size=10)
left_align_body(tbl, [1])
y += GAP_HEAD + SUB_TBL_H + GAP_SECTION

# 규제 참고
_txt(s, MARGIN_L, y, 25.26, LINE_H,
     "※ 유통산업발전법('12년) : 대형마트 영업 자정~오전 10시 제한·월 2회 의무휴업 → 전국 160여 PP센터 새벽배송 불가, 규제 완화 개정 시 새벽배송 확장 여지",
     size=9.5, bold=False, color=CRIMSON)

out = os.path.join(os.path.dirname(__file__), "ssg-quickcommerce-onepager.pptx")
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst), "| last y:", round(y, 2))
